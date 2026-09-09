from pathlib import Path
import re

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_ROW_HEIGHT_RULE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION_START
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / 'Panduan_Lengkap_IPSRS.md'
OUTPUT = ROOT / 'Panduan_Pengguna_IPSRS_Formal.docx'


def field(paragraph, instruction, cached_text):
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    begin.set(qn('w:dirty'), 'true')
    paragraph.add_run()._r.append(begin)
    code = OxmlElement('w:instrText')
    code.set(qn('xml:space'), 'preserve')
    code.text = instruction
    paragraph.add_run()._r.append(code)
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    paragraph.add_run()._r.append(separate)
    paragraph.add_run(cached_text)
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    paragraph.add_run()._r.append(end)


def inline(paragraph, text):
    replacements = {
        'Kehebatan Sistem (Mesin Pembelajar)': 'Penyimpanan Template Otomatis',
        'Kehebatan Sistem': 'Fungsi Otomatis Sistem',
        '⚠️ Peringatan Keras': 'Peringatan',
        '⚠️ Aturan Wajib': 'Ketentuan',
        'secara ajaib': 'secara otomatis',
        'dan voila!': 'dan',
        'keajaiban akan terjadi. ': '',
        'sistem akan diam-diam berlari ke gudang digital ini': 'sistem akan memperbarui catatan stok gudang',
        'langsung "dilempar" (diarahkan)': 'langsung diarahkan',
        'Kehebatan Sistem akan langsung bekerja': 'sistem akan memproses verifikasi',
        'Luar biasanya,': 'Selanjutnya,',
        'berteriak memberi peringatan': 'menampilkan peringatan',
        'Sistem sangat disiplin. ': '',
        'Tugas Anda beres!': 'Proses pemeliharaan selesai.',
        'merubah': 'mengubah',
    }
    for original, replacement in replacements.items():
        text = text.replace(original, replacement)
    for part in re.split(r'(\*\*.*?\*\*|\*[^*]+\*|`[^`]+`)', text):
        if part.startswith('**') and part.endswith('**'):
            paragraph.add_run(part[2:-2]).bold = True
        elif part.startswith('*') and part.endswith('*'):
            paragraph.add_run(part[1:-1]).italic = True
        elif part.startswith('`') and part.endswith('`'):
            paragraph.add_run(part[1:-1]).font.name = 'Consolas'
        else:
            paragraph.add_run(part)


def placeholder(document, caption):
    table = document.add_table(rows=1, cols=1)
    table.autofit = False
    table.columns[0].width = Cm(14)
    table.cell(0, 0).width = Cm(14)
    row = table.rows[0]
    row.height = Cm(6)
    row.height_rule = WD_ROW_HEIGHT_RULE.AT_LEAST
    row._tr.get_or_add_trPr().append(OxmlElement('w:cantSplit'))
    cell = row.cells[0]
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    properties = cell._tc.get_or_add_tcPr()
    shade = OxmlElement('w:shd')
    shade.set(qn('w:fill'), 'FFFFFF')
    properties.append(shade)
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        border = OxmlElement('w:' + edge)
        for key, value in [('val', 'single'), ('sz', '4'), ('color', '000000')]:
            border.set(qn('w:' + key), value)
        borders.append(border)
    properties.append(borders)
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.keep_with_next = True
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.add_run('[Sisipkan tangkapan layar di sini]').italic = True
    document.add_paragraph(caption, style='Caption')


def numbered_section(document, number_format):
    section = document.add_section(WD_SECTION_START.NEW_PAGE)
    section.different_first_page_header_footer = False
    section.header.is_linked_to_previous = False
    section.footer.is_linked_to_previous = False
    numbering = OxmlElement('w:pgNumType')
    numbering.set(qn('w:fmt'), number_format)
    numbering.set(qn('w:start'), '1')
    section._sectPr.append(numbering)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.first_line_indent = Cm(0)
    field(footer, ' PAGE ', 'i' if number_format == 'lowerRoman' else '1')


def build():
    content = SOURCE.read_text(encoding='utf-8-sig')
    body = content[content.index('Modul Manajemen Aset (Katalog & Inventaris Fisik)', content.index('<div')):]
    lines = [line.strip() for line in body.splitlines()]
    figure_chapters = {'1': '1', '2': '2', '2B': '3', '3': '4', '4': '5', '5': '6', '6': '7'}
    lines = [re.sub(r'^Gambar (2B|[1-6])\.', lambda match: 'Gambar ' + figure_chapters[match.group(1)] + '.', line) for line in lines]
    captions = [line for line in lines if line.startswith('Gambar ')]
    chapters = [line for line in lines if re.match(r'^Modul [A-Z]', line)]
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin = section.bottom_margin = Cm(3)
    section.left_margin, section.right_margin = Cm(4), Cm(3)
    section.header_distance = section.footer_distance = Cm(1.2)
    section.different_first_page_header_footer = True
    normal = document.styles['Normal']
    for style in document.styles:
        if style.type in (1, 2):
            style.font.name = 'Times New Roman'
            style.font.color.rgb = RGBColor(0, 0, 0)
    normal.font.name = 'Times New Roman'
    normal.font.size = Pt(12)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.first_line_indent = Cm(1.25)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.widow_control = True
    for name, size in [('Title', 16), ('Heading 1', 14), ('Heading 2', 12), ('Heading 3', 12)]:
        style = document.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.font.bold = True
        style.paragraph_format.first_line_indent = Cm(0)
        style.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.line_spacing = 1.15
        style.paragraph_format.keep_with_next = True
    document.styles['Heading 1'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.styles['Heading 1'].paragraph_format.space_after = Pt(18)
    document.styles['Title'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for name in ('Caption', 'List Bullet', 'List Number', 'Subtitle'):
        document.styles[name].paragraph_format.first_line_indent = Cm(0)
    document.styles['Caption'].font.size = Pt(11)
    document.styles['Caption'].font.italic = False
    document.styles['Caption'].paragraph_format.line_spacing = 1
    document.styles['Caption'].paragraph_format.space_before = Pt(6)
    document.styles['Caption'].paragraph_format.space_after = Pt(12)
    document.styles['Caption'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    settings = OxmlElement('w:updateFields')
    settings.set(qn('w:val'), 'true')
    document.settings.element.append(settings)
    document.core_properties.title = 'Panduan Pengguna Aplikasi Manajemen Aset IPSRS'
    document.core_properties.subject = 'Panduan operasional IPSRS RSUD Kota Yogyakarta'
    document.core_properties.author = ''

    for text, spacing, size, bold in [
        ('BUKU PANDUAN PENGGUNA', 18, 16, True),
        ('APLIKASI MANAJEMEN ASET IPSRS\nRSUD KOTA YOGYAKARTA', 60, 14, True),
        ('[LOGO INSTANSI]', 72, 12, False),
        ('Disusun oleh:\n[Nama penyusun]', 42, 12, False),
        ('Versi aplikasi: [Versi]\nTanggal penerbitan: [Tanggal]', 48, 12, False),
        ('INSTALASI PEMELIHARAAN SARANA RUMAH SAKIT\nRSUD KOTA YOGYAKARTA\n[TAHUN]', 0, 12, True),
    ]:
        paragraph = document.add_paragraph()
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.first_line_indent = Cm(0)
        paragraph.paragraph_format.space_after = Pt(spacing)
        paragraph.paragraph_format.line_spacing = 1.15
        run = paragraph.add_run(text)
        run.font.size = Pt(size)
        run.bold = bold

    numbered_section(document, 'lowerRoman')
    document.add_heading('Tentang Panduan Ini', level=1)
    document.add_paragraph('Buku ini membantu pengguna memahami pengelolaan sarana dan prasarana nonmedis IPSRS RSUD Kota Yogyakarta. Materi operasional bersumber dari Panduan_Lengkap_IPSRS.md [1], dengan acuan penyajian style_guide.md [2].')
    document.add_paragraph('Cakupannya meliputi katalog dan unit aset, QR Code dan lokasi, pelaporan kerusakan, pemeliharaan preventif, suku cadang, peminjaman, penghapusan, serta administrasi dan laporan.')
    document.add_heading('Cara Menggunakan Buku', level=2)
    for text in ['Administrator: gunakan seluruh modul sesuai kewenangan pengelolaan.', 'Teknisi: fokus pada penanganan laporan, pemeliharaan preventif, dan pencatatan penggunaan komponen.', 'Pelapor: gunakan Bab II untuk portal publik dan Bab III untuk akun pelapor internal.']:
        document.add_paragraph(text, style='List Bullet')
    document.add_paragraph('Materi disusun dalam tujuh bab. Modul pelapor internal disajikan tersendiri pada Bab III. Nomor gambar mengikuti bab tempat gambar dicantumkan. Istilah dan perilaku aplikasi dalam buku ini mengikuti dokumentasi sumber; kesesuaiannya dengan versi aplikasi yang dipakai perlu diperiksa sebelum penerbitan.')
    document.add_heading('Informasi Akses', level=2)
    document.add_paragraph('Alamat aplikasi: [Isi URL aplikasi]\nKontak bantuan IPSRS: [Isi kontak]\nPenanggung jawab dokumen: [Isi nama atau jabatan]')

    document.add_page_break()
    document.add_paragraph('DAFTAR ISI', style='Title')
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    field(paragraph, ' TOC \\o "1-2" \\h \\z \\u ', '\n'.join(['TENTANG PANDUAN INI'] + ['BAB ' + numeral + ' — ' + title[6:].upper() for numeral, title in zip(['I', 'II', 'III', 'IV', 'V', 'VI', 'VII'], chapters)] + ['DAFTAR PUSTAKA', 'LAMPIRAN — PETUNJUK PENYUNTINGAN']))
    document.add_paragraph('Perbarui field di Microsoft Word untuk menampilkan daftar isi lengkap beserta nomor halaman.')
    document.add_page_break()
    document.add_paragraph('DAFTAR GAMBAR', style='Title')
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.first_line_indent = Cm(0)
    paragraph.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    field(paragraph, ' TOC \\t "Caption,1" \\h \\z ', '\n'.join(captions))
    document.add_paragraph('Perbarui field di Microsoft Word setelah semua gambar disisipkan untuk menghitung nomor halaman.')

    chapter_number = 0
    subsection_number = 0
    roman = ['I', 'II', 'III', 'IV', 'V', 'VI', 'VII']
    for line in lines:
        if not line or line.startswith('<div') or line == '---':
            continue
        if line in chapters:
            if chapter_number == 0:
                numbered_section(document, 'decimal')
            else:
                document.add_page_break()
            document.add_heading('BAB ' + roman[chapter_number] + '\n' + line[6:].upper(), level=1)
            chapter_number += 1
            subsection_number = 0
        elif line.startswith('Gambar '):
            placeholder(document, line)
        elif line.startswith('Bagian '):
            document.add_paragraph(line, style='Heading 3')
        elif re.match(r'^\d+\. ', line) and ':' not in line:
            subsection_number += 1
            document.add_heading(f'{chapter_number}.{subsection_number} ' + re.sub(r'^\d+\. ', '', line), level=2)
        elif re.match(r'^[A-Z]\. ', line):
            paragraph = document.add_paragraph(style='Heading 3')
            inline(paragraph, line)
        elif line.startswith('- '):
            inline(document.add_paragraph(style='List Bullet'), line[2:])
        else:
            paragraph = document.add_paragraph()
            if line.endswith(':'):
                paragraph.paragraph_format.keep_with_next = True
            label = re.match(r'^([^:]{1,85}):\s*(.*)$', line)
            if label and '**' not in label[1]:
                paragraph.paragraph_format.first_line_indent = Cm(0)
                inline(paragraph, '**' + label[1] + ':** ')
                inline(paragraph, label[2])
            else:
                inline(paragraph, line)

    document.add_page_break()
    document.add_heading('Daftar Pustaka', level=1)
    document.add_paragraph('Daftar berikut mencantumkan dokumen internal yang digunakan dalam penyusunan buku. Identitas penulis dan tahun tidak tercantum pada sumber.')
    document.add_paragraph('[1] “Panduan Lengkap IPSRS.” Tanpa tahun. Dokumen internal proyek, docs/Panduan_Lengkap_IPSRS.md.')
    document.add_paragraph('[2] “Panduan Gaya Penulisan (Style Guide) User Guide Book IPSRS.” Tanpa tahun. Dokumen internal proyek, docs/style_guide.md.')
    document.add_paragraph('Catatan penyunting: lengkapi identitas penulis, tahun, dan referensi SOP atau dokumen resmi yang benar-benar digunakan sebelum buku diterbitkan.')
    document.add_page_break()
    document.add_heading('Lampiran — Petunjuk Penyuntingan', level=1)
    for text in ['Lengkapi identitas penyusun, logo, versi, tanggal penerbitan, alamat aplikasi, dan kontak bantuan.', 'Pada setiap kotak gambar, pilih teks petunjuk di dalam sel lalu hapus. Sisipkan screenshot melalui Insert > Pictures. Gunakan In Line with Text dan lebar maksimal 13 cm agar gambar mengikuti tata letak halaman.', 'Pertahankan paragraf caption di bawah kotak. Daftar gambar mengambil semua paragraf bergaya Caption. Penomoran caption mengikuti naskah sumber; jika menambah atau memindahkan gambar, sesuaikan nomornya.', 'Di Microsoft Word, tekan Ctrl+A lalu F9. Jika diminta, pilih Update entire table. Lakukan pada daftar isi dan daftar gambar; periksa juga nomor halaman footer. Field belum dihitung oleh generator dokumen.', 'Periksa nama menu, tombol, izin peran, dan langkah operasional terhadap aplikasi yang digunakan. Tambahkan referensi resmi hanya jika tersedia dan digunakan.', 'Setelah gambar, daftar, dan tata letak selesai diperiksa, hapus catatan penyunting serta lampiran ini jika tidak diperlukan dalam edisi pengguna.']:
        document.add_paragraph(text, style='List Number')
    for paragraph in document.paragraphs:
        if paragraph.style.name == 'Heading 1' and not paragraph.text.startswith('BAB '):
            for run in paragraph.runs:
                run.text = run.text.upper()
        if paragraph.text.startswith('[1]') or paragraph.text.startswith('[2]'):
            paragraph.paragraph_format.left_indent = Cm(1)
            paragraph.paragraph_format.first_line_indent = Cm(-1)
            paragraph.paragraph_format.line_spacing = 1
    document.save(OUTPUT)
    print(f'Created: {OUTPUT}\nChapters: {len(chapters)}\nImage placeholders: {len(captions)}')


if __name__ == '__main__':
    build()
